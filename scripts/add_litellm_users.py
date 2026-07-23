#!/usr/bin/env python3
"""scripts/add_litellm_users.py — LiteLLM 프록시 사용자 계정 일괄 등록.

엑셀 파일(기본 scripts/users.xlsx)에 채워진 행을 읽어 LiteLLM Admin API(POST /user/new,
POST /team/member_add)로 계정을 만든다. Postgres/Prisma DB에 직접 쓰지 않는 이유: 가상 키
해시·budget id 등 litellm 내부 생성 규칙을 손으로 재현하는 것보다 API 호출이 litellm 버전이
올라가도 덜 깨진다(2026-07-23 결정, gr-lgcaip-project-fab0704-backend는 STORE_MODEL_IN_DB=False로
이 DB를 가상 키/팀/유저 전용으로 쓰는 LiteLLM 프록시 게이트웨이 — config.yaml 참고).

`auto_create_key`를 명시적으로 False로 보낸다 — /user/new는 기본값이 True라 그냥 두면 사용자당
가상 API 키가 자동 발급된다(이번 용도에서는 계정만 필요해 명시적으로 제외).

비밀번호는 /user/new가 아니라 /user/update로 별도 설정한다 — litellm 1.93.0의 NewUserRequest
스키마에는 password 필드 자체가 없고(UpdateUserRequest에만 존재), 서버가 이 값을 해싱해 저장한다
(설치된 litellm 1.93.0 소스로 직접 확인, 2026-07-23). 엑셀 password 컬럼이 비어 있는 신규
사용자는 DEFAULT_PASSWORD(사용자 요청으로 "1234")를 사용한다 — 최초 로그인 후 변경을 안내할 것.
기존 사용자는 password 컬럼에 값이 있을 때만 비밀번호를 변경한다(빈 값이면 건드리지 않음).

CI/CD에 엮여 있지 않다 — 사람이 로컬에서 수동으로 실행하는 용도(먼저 --dry-run으로 확인 권장).
필드명(user_role enum 값, /team/member_add의 정확한 body 형태)은 litellm 1.93.0 실서버로
라이브 검증되지 않았다(사내망 서버라 이 스크립트를 작성한 환경에서 접근 불가) — 실제 응답이
다르면 아래 DEFAULT_* 상수와 API 함수(create_user/add_team_member)만 고치면 된다. 서버가
FastAPI라 <base_url>/docs(Swagger)에서 정확한 스키마를 대조할 수 있다.
"""

from __future__ import annotations

import argparse
import os
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

import requests
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook

DEFAULT_USER_ROLE = "internal_user"
DEFAULT_TEAM_MEMBER_ROLE = "user"
DEFAULT_PASSWORD = "1234"

INPUT_COLUMNS = ["user_email", "user_id", "user_role", "team_id", "password"]
STATUS_COLUMNS = ["status", "message", "processed_at"]
STATUS_DONE = "완료"
STATUS_FAILED = "실패"

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_XLSX = SCRIPT_DIR / "users.xlsx"
DEFAULT_ENV_FILE = SCRIPT_DIR / ".env"


@dataclass
class Config:
    base_url: str
    master_key: str


class LiteLLMApiError(Exception):
    """LiteLLM Admin API가 non-2xx를 반환했을 때 — 원문 응답 본문을 그대로 보존한다(엑셀 message
    컬럼에 그대로 적어야 사용자가 필드명 불일치 등 원인을 직접 진단할 수 있다)."""

    def __init__(self, method: str, url: str, status_code: int, body: str):
        self.method = method
        self.url = url
        self.status_code = status_code
        self.body = body
        super().__init__(f"{method} {url} -> {status_code}: {body}")


class AuthError(Exception):
    """마스터키 오설정 등 401/403 — 행 전체를 실패 처리하기 전에 실행 자체를 중단시키는 신호."""


def load_config(env_file: Path) -> Config:
    if env_file.exists():
        load_dotenv(env_file)
    base_url = os.environ.get("LITELLM_BASE_URL", "").strip()
    master_key = os.environ.get("LITELLM_MASTER_KEY", "").strip()
    if not base_url or not master_key:
        raise SystemExit(
            "[설정 오류] LITELLM_BASE_URL / LITELLM_MASTER_KEY가 비어 있습니다.\n"
            f"'{env_file}' 파일을 만들고 '{SCRIPT_DIR / '.env.example'}'을 참고해 값을 채워주세요."
        )
    return Config(base_url=base_url.rstrip("/"), master_key=master_key)


def build_session(config: Config) -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "Authorization": f"Bearer {config.master_key}",
            "Content-Type": "application/json",
        }
    )
    return session


def _request(session: requests.Session, method: str, base_url: str, path: str, **kwargs):
    return session.request(method, f"{base_url}{path}", timeout=30, **kwargs)


def get_user_info(session: requests.Session, base_url: str, user_id: str) -> dict | None:
    """GET /user/info — 200이면 기존 유저 정보, 401/403이면 AuthError, 그 외 non-2xx는
    '미존재'로 취급(fail-open — litellm 버전별 404/400 등 표현 차이를 굳이 다 맞추지 않는다)."""
    resp = _request(session, "GET", base_url, "/user/info", params={"user_id": user_id})
    if resp.status_code == 200:
        return resp.json()
    if resp.status_code in (401, 403):
        raise AuthError(f"GET /user/info -> {resp.status_code}: {resp.text}")
    return None


def create_user(session: requests.Session, base_url: str, *, user_id: str, user_email: str, user_role: str) -> dict:
    payload = {
        "user_id": user_id,
        "user_email": user_email,
        "user_role": user_role,
        "auto_create_key": False,
    }
    resp = _request(session, "POST", base_url, "/user/new", json=payload)
    if resp.status_code in (401, 403):
        raise AuthError(f"POST /user/new -> {resp.status_code}: {resp.text}")
    if resp.status_code >= 300:
        raise LiteLLMApiError("POST", resp.url, resp.status_code, resp.text)
    return resp.json()


def add_team_member(
    session: requests.Session,
    base_url: str,
    *,
    team_id: str,
    user_id: str,
    user_email: str,
    role: str = DEFAULT_TEAM_MEMBER_ROLE,
) -> dict:
    payload = {
        "team_id": team_id,
        "member": {"user_id": user_id, "user_email": user_email, "role": role},
    }
    resp = _request(session, "POST", base_url, "/team/member_add", json=payload)
    if resp.status_code in (401, 403):
        raise AuthError(f"POST /team/member_add -> {resp.status_code}: {resp.text}")
    if resp.status_code >= 300:
        raise LiteLLMApiError("POST", resp.url, resp.status_code, resp.text)
    return resp.json()


def update_user_password(session: requests.Session, base_url: str, *, user_id: str, password: str) -> dict:
    payload = {"user_id": user_id, "password": password}
    resp = _request(session, "POST", base_url, "/user/update", json=payload)
    if resp.status_code in (401, 403):
        raise AuthError(f"POST /user/update -> {resp.status_code}: {resp.text}")
    if resp.status_code >= 300:
        raise LiteLLMApiError("POST", resp.url, resp.status_code, resp.text)
    return resp.json()


def derive_user_id(user_email: str) -> str:
    return user_email.split("@", 1)[0]


def ensure_columns(ws, columns: list[str]) -> dict[str, int]:
    """헤더 행(1행)에서 컬럼명 -> 1-based 인덱스 매핑을 만든다. 없는 컬럼(status/message/
    processed_at)은 헤더 오른쪽 끝에 이어 붙인다 — 사용자가 손으로 만든 시트도 그대로 동작."""
    header = {}
    max_col = ws.max_column or 0
    for col in range(1, max_col + 1):
        value = ws.cell(row=1, column=col).value
        if value:
            header[str(value).strip()] = col
    for name in columns:
        if name not in header:
            max_col += 1
            ws.cell(row=1, column=max_col, value=name)
            header[name] = max_col
    return header


def iter_data_rows(ws):
    for row in range(2, (ws.max_row or 1) + 1):
        if all(ws.cell(row=row, column=c).value in (None, "") for c in range(1, (ws.max_column or 1) + 1)):
            continue
        yield row


def now_iso() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


@dataclass
class RowResult:
    row: int
    user_email: str
    ok: bool
    message: str


def process_row(session: requests.Session, base_url: str, ws, header: dict[str, int], row: int, *, dry_run: bool) -> RowResult:
    def get(col_name: str) -> str:
        value = ws.cell(row=row, column=header[col_name]).value
        return str(value).strip() if value not in (None, "") else ""

    status = get("status")
    user_email = get("user_email")

    if status == STATUS_DONE:
        return RowResult(row, user_email, True, "이미 완료 - 스킵")

    if not user_email:
        return RowResult(row, user_email, False, "user_email 필수")

    user_id = get("user_id") or derive_user_id(user_email)
    user_role = get("user_role") or DEFAULT_USER_ROLE
    team_ids = [t.strip() for t in get("team_id").split(",") if t.strip()]
    password_input = get("password")

    exists = get_user_info(session, base_url, user_id) is not None
    # 신규 사용자는 항상 비밀번호를 설정한다(입력값 없으면 DEFAULT_PASSWORD). 기존 사용자는 password
    # 컬럼에 값이 있을 때만 변경 - 빈 값으로 왔다고 기존 비밀번호를 건드리지 않는다.
    password_to_set = password_input or (None if exists else DEFAULT_PASSWORD)

    if dry_run:
        action = "이미 존재 - team만 처리 예정" if exists else "신규 생성 예정"
        team_note = f", team={team_ids}" if team_ids else ""
        if password_to_set is None:
            pw_note = ""
        elif exists:
            pw_note = ", 비밀번호 변경 예정(입력값)"
        else:
            pw_note = f", 비밀번호=입력값" if password_input else f", 비밀번호={DEFAULT_PASSWORD}(기본값)"
        return RowResult(row, user_email, True, f"[dry-run] {action} (user_id={user_id}, role={user_role}{team_note}{pw_note})")

    if not exists:
        try:
            create_user(session, base_url, user_id=user_id, user_email=user_email, user_role=user_role)
        except LiteLLMApiError as exc:
            return RowResult(row, user_email, False, str(exc))

    pw_note = ""
    if password_to_set is not None:
        try:
            update_user_password(session, base_url, user_id=user_id, password=password_to_set)
        except LiteLLMApiError as exc:
            return RowResult(row, user_email, False, f"비밀번호 설정 실패 - {exc}")
        if password_input:
            pw_note = ", 비밀번호 변경됨" if exists else ", 비밀번호 설정됨(입력값)"
        else:
            pw_note = f", 비밀번호 설정됨(기본값 {DEFAULT_PASSWORD} - 최초 로그인 후 변경 안내 필요)"

    team_errors = []
    for team_id in team_ids:
        try:
            add_team_member(session, base_url, team_id=team_id, user_id=user_id, user_email=user_email)
        except LiteLLMApiError as exc:
            team_errors.append(f"{team_id}: {exc}")

    if team_errors:
        return RowResult(row, user_email, False, "team 추가 실패 - " + "; ".join(team_errors))

    prefix = "이미 존재 - team 추가 완료" if exists else "생성 완료"
    team_note = f" (team: {', '.join(team_ids)})" if team_ids else ""
    return RowResult(row, user_email, True, f"{prefix}{team_note}{pw_note}")


def write_result(ws, header: dict[str, int], row: int, result: RowResult) -> None:
    ws.cell(row=row, column=header["status"], value=STATUS_DONE if result.ok else STATUS_FAILED)
    ws.cell(row=row, column=header["message"], value=result.message)
    ws.cell(row=row, column=header["processed_at"], value=now_iso())


def init_workbook(path: Path) -> None:
    if path.exists():
        raise SystemExit(f"[오류] 이미 존재하는 파일입니다: {path}")
    wb = Workbook()
    ws = wb.active
    for i, name in enumerate(INPUT_COLUMNS, start=1):
        ws.cell(row=1, column=i, value=name)
    wb.save(path)
    print(f"헤더만 있는 빈 파일을 만들었습니다: {path}")


def run(config: Config, file_path: Path, sheet: str | None, dry_run: bool) -> int:
    if not file_path.exists():
        raise SystemExit(f"[오류] 파일이 없습니다: {file_path} (--init으로 먼저 만드세요)")

    wb = load_workbook(file_path)
    ws = wb[sheet] if sheet else wb.active
    header = ensure_columns(ws, INPUT_COLUMNS + STATUS_COLUMNS)
    missing = [c for c in ["user_email"] if c not in header]
    if missing:
        raise SystemExit(f"[오류] 필수 컬럼이 없습니다: {missing}")

    session = build_session(config)
    failed = 0
    processed = 0

    for row in iter_data_rows(ws):
        try:
            result = process_row(session, config.base_url, ws, header, row, dry_run=dry_run)
        except AuthError as exc:
            print(f"[인증 오류] {exc}")
            print("마스터키 또는 서버 주소를 확인하세요 — 나머지 행은 처리하지 않고 중단합니다.")
            if not dry_run:
                wb.save(file_path)
            return 1

        mark = "OK" if result.ok else "FAIL"
        print(f"[{mark}] row {result.row} ({result.user_email}): {result.message}")
        if not result.ok:
            failed += 1
        processed += 1

        if not dry_run:
            write_result(ws, header, row, result)
            wb.save(file_path)

    print(f"\n총 {processed}건 처리, 실패 {failed}건." + (" (dry-run — 저장 안 함)" if dry_run else ""))
    return 1 if failed else 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", type=Path, default=DEFAULT_XLSX, help=f"엑셀 파일 경로 (기본: {DEFAULT_XLSX})")
    parser.add_argument("--env-file", type=Path, default=DEFAULT_ENV_FILE, help=f"설정 파일 경로 (기본: {DEFAULT_ENV_FILE})")
    parser.add_argument("--sheet", type=str, default=None, help="시트 이름 (기본: 활성 시트)")
    parser.add_argument("--dry-run", action="store_true", help="API 조회만 하고 실제 생성/저장은 하지 않음")
    parser.add_argument("--init", action="store_true", help="--file 경로에 헤더만 있는 빈 엑셀을 생성하고 종료")
    args = parser.parse_args(argv)

    if args.init:
        init_workbook(args.file)
        return 0

    config = load_config(args.env_file)
    return run(config, args.file, args.sheet, args.dry_run)


if __name__ == "__main__":
    sys.exit(main())
